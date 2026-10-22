# -*- coding: utf-8 -*-
"""67 · Markdown → Word（.docx）转换器，给交接/技术文档用。⭐零 GPU、纯 CPU、秒级。

═══════════════════════════════════════════════════════════════════════
为什么自己写而不用 pandoc
═══════════════════════════════════════════════════════════════════════
本机没有 pandoc（2026-09-17 实测），但 `D:\Anaconda\envs\dl` 里有 python-docx。
交接文档要给非技术读者用 Word 看，而 Markdown 源要留给后续会话继续编辑 ——
所以走「Markdown 是源、docx 是导出物」的路：源改了重跑一次即可，⛔别在 Word 里直接改。

支持的 Markdown 子集（够本项目所有文档用）：
  # ~ ####  标题        - / * / 1.  列表       | a | b |  表格（含 <br>、\\| 转义）
  ``` 代码块            > 引用块（内部可再有 #/| 结构，会被剥掉 > 后按普通行处理）
  **粗体**  `行内代码`  ~~删除线~~            ---  水平线（渲染为细分隔）
  YAML front matter（--- title: --- 开头）→ 取 title 当封面标题

用法：
  D:\\Anaconda\\envs\\dl\\python.exe scripts\\67_md2docx.py `
     --md "docs\\交接日志\\AIC-3714_步骤记录与提分规划_2026-09-18.md" `
     --out "docs\\交接日志\\AIC-3714_步骤记录与提分规划_2026-09-18.docx" `
     --copy-to "D:\\aic_data\\chat_docs\\【5 交接】步骤记录与提分规划_2026-09-18.docx"
"""
import argparse
import os
import re
import shutil
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor, Cm

BODY_FONT = "Microsoft YaHei"      # 中英文都能显示；Word 里东亚字体单独设
CODE_FONT = "Consolas"
CODE_BG = "F2F2F2"
QUOTE_COLOR = RGBColor(0x55, 0x55, 0x55)

INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`|~~.+?~~)")


# ══════════════════════════════════════════════════════════ 底层：字体 / 底纹
def _set_run_font(run, name=BODY_FONT, size=None, bold=None, italic=None,
                  color=None, strike=None, mono=False):
    f = run.font
    f.name = CODE_FONT if mono else name
    rpr = run._element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    rfonts.set(qn("w:eastAsia"), CODE_FONT if mono else name)
    rfonts.set(qn("w:ascii"), CODE_FONT if mono else name)
    rfonts.set(qn("w:hAnsi"), CODE_FONT if mono else name)
    if size is not None:
        f.size = Pt(size)
    if bold is not None:
        f.bold = bold
    if italic is not None:
        f.italic = italic
    if color is not None:
        f.color.rgb = color
    if strike:
        f.strike = True


def _shade(paragraph_or_cell, hex_fill):
    """给段落或单元格加底纹。"""
    el = paragraph_or_cell._p if hasattr(paragraph_or_cell, "_p") else paragraph_or_cell._tc
    pr = el.get_or_add_pPr() if hasattr(el, "get_or_add_pPr") else el.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    pr.append(shd)


def add_inline(paragraph, text, size=None, base_bold=False, color=None):
    """把带 **粗体** / `代码` / ~~删除线~~ 的一行文本写成多个 run。"""
    text = text.replace("\\|", "|")
    for piece in INLINE_RE.split(text):
        if not piece:
            continue
        if piece.startswith("**") and piece.endswith("**") and len(piece) > 4:
            r = paragraph.add_run(piece[2:-2])
            _set_run_font(r, size=size, bold=True, color=color)
        elif piece.startswith("`") and piece.endswith("`") and len(piece) > 2:
            r = paragraph.add_run(piece[1:-1])
            _set_run_font(r, size=(size or 10.5) - 0.5, mono=True, color=RGBColor(0x9C, 0x27, 0x06))
        elif piece.startswith("~~") and piece.endswith("~~") and len(piece) > 4:
            r = paragraph.add_run(piece[2:-2])
            _set_run_font(r, size=size, strike=True, color=RGBColor(0x88, 0x88, 0x88))
        else:
            r = paragraph.add_run(piece)
            _set_run_font(r, size=size, bold=(True if base_bold else None), color=color)


# ══════════════════════════════════════════════════════════ 块级元素
def add_heading(doc, text, level):
    h = doc.add_heading(level=min(level, 4))
    add_inline(h, text.strip(), size={1: 18, 2: 15, 3: 13, 4: 12}.get(level, 12), base_bold=True,
               color=RGBColor(0x1F, 0x3A, 0x5F))
    return h


def add_paragraph(doc, text, quote=False):
    p = doc.add_paragraph()
    if quote:
        p.paragraph_format.left_indent = Cm(0.8)
        add_inline(p, text, size=10.5, color=QUOTE_COLOR)
        _shade(p, "F7F7F7")
    else:
        add_inline(p, text, size=10.5)
    p.paragraph_format.space_after = Pt(4)
    return p


def add_list_item(doc, text, numbered=False, level=0):
    style = "List Number" if numbered else "List Bullet"
    if level > 0:
        style += " %d" % min(level + 1, 3)
    try:
        p = doc.add_paragraph(style=style)
    except KeyError:
        p = doc.add_paragraph(style="List Bullet")
    add_inline(p, text, size=10.5)
    p.paragraph_format.space_after = Pt(2)
    return p


def add_code_block(doc, lines):
    for ln in lines:
        p = doc.add_paragraph()
        r = p.add_run(ln if ln else " ")
        _set_run_font(r, size=9, mono=True)
        p.paragraph_format.space_after = Pt(0)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.left_indent = Cm(0.5)
        _shade(p, CODE_BG)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def split_row(line):
    """按未转义的 | 切表格行；单元格里的 \\| 保留成 |。"""
    line = line.strip()
    if line.startswith("|"):
        line = line[1:]
    if line.endswith("|") and not line.endswith("\\|"):
        line = line[:-1]
    cells, buf, i = [], "", 0
    while i < len(line):
        c = line[i]
        if c == "\\" and i + 1 < len(line) and line[i + 1] == "|":
            buf += "|"
            i += 2
            continue
        if c == "|":
            cells.append(buf.strip())
            buf = ""
        else:
            buf += c
        i += 1
    cells.append(buf.strip())
    return cells


def is_sep_row(cells):
    return all(re.fullmatch(r":?-{2,}:?", c.strip()) for c in cells if c.strip()) and any(cells)


def add_table(doc, rows):
    header, body = rows[0], rows[1:]
    ncol = len(header)
    body = [r + [""] * (ncol - len(r)) if len(r) < ncol else r[:ncol] for r in body]
    t = doc.add_table(rows=1 + len(body), cols=ncol)
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for j, cell_text in enumerate(header):
        cell = t.rows[0].cells[j]
        cell.text = ""
        add_inline(cell.paragraphs[0], cell_text, size=9.5, base_bold=True)
        _shade(cell, "DCE6F1")
    for i, row in enumerate(body, start=1):
        for j, cell_text in enumerate(row):
            cell = t.rows[i].cells[j]
            cell.text = ""
            parts = re.split(r"<br\s*/?>", cell_text)
            for k, part in enumerate(parts):
                para = cell.paragraphs[0] if k == 0 else cell.add_paragraph()
                add_inline(para, part.strip(), size=9.5)
                para.paragraph_format.space_after = Pt(0)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)


def add_hr(doc):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(2)
    pPr = p._p.get_or_add_pPr()
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "BBBBBB")
    pbdr.append(bottom)
    pPr.append(pbdr)


# ══════════════════════════════════════════════════════════ 主转换
def convert(md_path, out_path):
    with open(md_path, encoding="utf-8") as f:
        lines = f.read().splitlines()

    doc = Document()
    # 页边距与默认字体
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(2.0)
        s.top_margin = s.bottom_margin = Cm(2.0)
    normal = doc.styles["Normal"]
    normal.font.name = BODY_FONT
    normal.font.size = Pt(10.5)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), BODY_FONT)

    i, n = 0, len(lines)
    title = None
    # YAML front matter
    if i < n and lines[i].strip() == "---":
        j = i + 1
        while j < n and lines[j].strip() != "---":
            m = re.match(r'\s*title:\s*"?(.*?)"?\s*$', lines[j])
            if m:
                title = m.group(1)
            j += 1
        i = j + 1
    if title:
        tp = doc.add_paragraph()
        tp.alignment = WD_ALIGN_PARAGRAPH.CENTER
        add_inline(tp, title, size=20, base_bold=True, color=RGBColor(0x1F, 0x3A, 0x5F))

    para_buf, quote_mode = [], False
    stats = {"heading": 0, "para": 0, "table": 0, "code": 0, "list": 0}

    def flush_para():
        nonlocal para_buf
        if para_buf:
            add_paragraph(doc, " ".join(x.strip() for x in para_buf), quote=quote_mode)
            stats["para"] += 1
            para_buf = []

    while i < n:
        raw = lines[i]
        line = raw.rstrip("\n")
        stripped = line.strip()

        # 引用块：剥掉 "> " 后按普通行处理（保留 quote_mode 给段落上底纹）
        if stripped.startswith(">"):
            quote_mode = True
            line = re.sub(r"^\s*>\s?", "", line)
            stripped = line.strip()
            if stripped == "":
                flush_para()
                i += 1
                continue
        else:
            if quote_mode and stripped != "" and not stripped.startswith(">"):
                flush_para()
                quote_mode = False

        # 代码块
        if stripped.startswith("```"):
            flush_para()
            j = i + 1
            block = []
            while j < n and not lines[j].strip().startswith("```"):
                block.append(re.sub(r"^\s*>\s?", "", lines[j]) if quote_mode else lines[j])
                j += 1
            add_code_block(doc, block)
            stats["code"] += 1
            i = j + 1
            continue

        # 表格
        if stripped.startswith("|"):
            flush_para()
            rows = []
            while i < n:
                l2 = re.sub(r"^\s*>\s?", "", lines[i]) if quote_mode else lines[i]
                if not l2.strip().startswith("|"):
                    break
                cells = split_row(l2)
                if not is_sep_row(cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
                stats["table"] += 1
            continue

        # 空行
        if stripped == "":
            flush_para()
            i += 1
            continue

        # 水平线
        if re.fullmatch(r"-{3,}|\*{3,}", stripped):
            flush_para()
            add_hr(doc)
            i += 1
            continue

        # 标题
        m = re.match(r"^(#{1,6})\s+(.*)$", stripped)
        if m:
            flush_para()
            add_heading(doc, m.group(2), len(m.group(1)))
            stats["heading"] += 1
            i += 1
            continue

        # 列表
        m = re.match(r"^(\s*)([-*]|\d+[.)])\s+(.*)$", line)
        if m:
            flush_para()
            indent = len(m.group(1).replace("\t", "    "))
            numbered = m.group(2)[0].isdigit()
            text = m.group(3)
            text = re.sub(r"^\[( |x|X)\]\s*", lambda mm: ("☐ " if mm.group(1) == " " else "☑ "), text)
            add_list_item(doc, text, numbered=numbered, level=indent // 2)
            stats["list"] += 1
            i += 1
            continue

        # 普通段落（连续行合并）
        para_buf.append(line)
        i += 1

    flush_para()
    doc.save(out_path)
    return stats


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("--md", required=True)
    ap.add_argument("--out", required=True)
    ap.add_argument("--copy-to", default=None, help="再复制一份到这个路径（如 chat_docs）")
    args = ap.parse_args()
    stats = convert(args.md, args.out)
    size = os.path.getsize(args.out)
    print("✅ 已写出 %s  (%.0f KB)" % (args.out, size / 1024))
    print("   标题 %d / 段落 %d / 表格 %d / 代码块 %d / 列表项 %d"
          % (stats["heading"], stats["para"], stats["table"], stats["code"], stats["list"]))
    if args.copy_to:
        os.makedirs(os.path.dirname(args.copy_to), exist_ok=True)
        shutil.copy2(args.out, args.copy_to)
        print("   已复制到 %s" % args.copy_to)


if __name__ == "__main__":
    main()
